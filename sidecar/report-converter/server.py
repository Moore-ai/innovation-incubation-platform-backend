import os, re, subprocess, tempfile, uuid
from fastapi import FastAPI, HTTPException
from pydantic import BaseModel

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_MD2PDF = os.path.join(_SCRIPT_DIR, "venv", "Scripts", "md2pdf.exe")

app = FastAPI()

class ConvertRequest(BaseModel):
    markdown: str
    title: str = ""
    font_path: str = ""

def _render_mermaid(md: str, out_dir: str) -> str:
    """用 md2pdf 的渲染器将 Mermaid 块转为 PNG，替换为 ![]() 引用。"""
    from md2pdf import render_mermaid_to_png
    pattern = r"```mermaid\n(.*?)```"
    def _repl(m):
        code = m.group(1).strip()
        if not code:
            return ""
        png = os.path.join(out_dir, f"chart_{uuid.uuid4().hex}.png")
        try:
            render_mermaid_to_png(code, png, width=800, height=600, scale=2)
            if os.path.exists(png):
                return f"![chart]({png})"
        except Exception:
            pass
        return f"```\n{code}\n```"
    return re.sub(pattern, _repl, md, flags=re.DOTALL)

@app.post("/convert/pdf")
def convert_pdf(req: ConvertRequest):
    try:
        workdir = tempfile.mkdtemp()
        md_path = os.path.join(workdir, "report.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(req.markdown)
        output = os.path.join(workdir, f"report_{uuid.uuid4().hex}.pdf")
        font = req.font_path or r"C:\Windows\Fonts\simhei.ttf"
        subprocess.run([_MD2PDF, md_path, "-o", output,
                        "--font", font,
                        "--mermaid-scale", "2"],
                      capture_output=True, timeout=120, check=True)
        return {"file_path": output}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

def _add_formatted_paragraph(doc, text: str):
    """解析 inline Markdown 格式（**粗体**, *斜体*）写入 docx 段落。"""
    para = doc.add_paragraph()
    token_re = re.compile(r"(\*\*([^*]+?)\*\*|\*([^*]+?)\*)")
    idx = 0
    for m in token_re.finditer(text):
        if m.start() > idx:
            para.add_run(text[idx:m.start()])
        if m.group(2):
            para.add_run(m.group(2)).bold = True
        elif m.group(3):
            para.add_run(m.group(3)).italic = True
        idx = m.end()
    if idx < len(text):
        para.add_run(text[idx:])
    return para


@app.post("/convert/docx")
def convert_docx(req: ConvertRequest):
    try:
        from docx import Document
        from docx.shared import Inches
        workdir = tempfile.mkdtemp()
        md = _render_mermaid(req.markdown, workdir)
        doc = Document()
        current_table = None
        img_count = 0
        for line in md.split("\n"):
            if not line.strip():
                current_table = None
                continue
            # Track images
            if line.startswith("![chart]("):
                img_count += 1
            if line.startswith("# ") and not line.startswith("## "):
                doc.add_heading(line[2:].strip(), 0)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), 1)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), 2)
            elif line.startswith("|"):
                cells = [c.strip() for c in line.split("|")[1:-1]]
                if cells:
                    if current_table is None:
                        current_table = doc.add_table(rows=1, cols=len(cells))
                        current_table.style = "Table Grid"
                        for j, c in enumerate(cells):
                            current_table.rows[0].cells[j].text = c
                    else:
                        row = current_table.add_row()
                        for j, c in enumerate(cells):
                            if j < len(row.cells):
                                row.cells[j].text = c
            elif line.startswith("![chart]("):
                img_path = line[line.index("(")+1:line.index(")")]
                if os.path.exists(img_path):
                    doc.add_picture(img_path, width=Inches(5.5))
                    doc.add_paragraph("")
            elif line.startswith("---"):
                continue
            elif line.startswith("```"):
                continue
            else:
                text = line.strip()
                if text and not text.startswith("!["):
                    _add_formatted_paragraph(doc, text)
        output = os.path.join(workdir, f"report_{uuid.uuid4().hex}.docx")
        doc.save(output)
        return {"file_path": output}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

if __name__ == "__main__":
    import uvicorn
    try: 
        uvicorn.run(app, host="127.0.0.1", port=9800)
    except KeyboardInterrupt:
        pass
